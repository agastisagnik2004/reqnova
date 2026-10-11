"""Document ingestion: turn an uploaded file into plain text.

Routing (traditional code decides, the vision model only does what code can't):
  .txt / .md          -> read directly
  .docx               -> python-docx (paragraphs + tables)
  .pdf                -> PyMuPDF text layer; pages with almost no text
                         (scans, screenshots) are rendered and sent to the vision model
  .png / .jpg / .webp -> vision model (OCR + table structure -> markdown)
"""
from __future__ import annotations

import base64
import io
import json
import re
import logging

import requests
from PIL import Image

from . import config

log = logging.getLogger("parser")
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
_URL = f"{config.API_BASE_URL}/chat/completions"
_MAX_SIDE = 1600  # the vision model works best on page-sized images


class ParseError(RuntimeError):
    pass


def _to_b64_jpeg(img: Image.Image) -> str:
    img = img.convert("RGB")
    if max(img.size) > _MAX_SIDE:
        img.thumbnail((_MAX_SIDE, _MAX_SIDE))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=88)
    return base64.b64encode(buf.getvalue()).decode()


def _headers() -> dict:
    return {"Authorization": f"Bearer {config.VISION_API_KEY}", "Accept": "application/json"}


def _parse_v2(b64: str) -> str:
    """Primary vision model: OpenAI-style image_url message."""
    body = {"model": config.VISION_MODEL, "max_tokens": 3000, "temperature": 0,
            "messages": [{"role": "user", "content": [
                {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}"}}]}]}
    r = requests.post(_URL, headers=_headers(), json=body, timeout=180)
    r.raise_for_status()
    msg = r.json()["choices"][0]["message"]
    if msg.get("content"):
        return msg["content"]
    for tc in msg.get("tool_calls") or []:
        return _tool_args_to_text(tc["function"]["arguments"])
    raise ParseError("empty response from parse model")


def _parse_v1(b64: str) -> str:
    """Fallback vision model: <img> tag + forced `markdown_no_bbox` tool call."""
    tool = "markdown_no_bbox"
    body = {"model": config.VISION_FALLBACK_MODEL, "max_tokens": 3500,
            "messages": [{"role": "user", "content": f'<img src="data:image/jpeg;base64,{b64}" />'}],
            "tools": [{"type": "function", "function": {"name": tool}}],
            "tool_choice": {"type": "function", "function": {"name": tool}}}
    r = requests.post(_URL, headers=_headers(), json=body, timeout=180)
    r.raise_for_status()
    calls = r.json()["choices"][0]["message"].get("tool_calls") or []
    if not calls:
        raise ParseError("parse model returned no tool call")
    return _tool_args_to_text(calls[0]["function"]["arguments"])


def _tool_args_to_text(args: str) -> str:
    data = json.loads(args)
    if isinstance(data, list):  # [{"text": ...}] or [[{...}]]
        flat = data[0] if data and isinstance(data[0], list) else data
        return "\n".join(d.get("text", "") for d in flat if isinstance(d, dict))
    return str(data)


def latex_tables_to_markdown(text: str) -> str:
    """The vision model returns tables as LaTeX `tabular`; convert to markdown rows."""
    def conv(m: re.Match) -> str:
        rows = [r.strip() for r in m.group(1).split("\\\\") if r.strip() and not r.strip().startswith("\\hline")]
        lines = []
        for r in rows:
            r = r.replace("\\hline", "").strip()
            cells = [c.strip() for c in r.split("&")]
            lines.append("| " + " | ".join(cells) + " |")
        return "\n".join(lines)
    return re.sub(r"\\begin\{tabular\}\{[^}]*\}(.*?)\\end\{tabular\}", conv, text, flags=re.S)


def parse_image(img: Image.Image) -> tuple[str, str]:
    """Return (markdown_text, model_used). Tries 2.0, then v1."""
    b64 = _to_b64_jpeg(img)
    errors = []
    for fn, name in ((_parse_v2, "vision"), (_parse_v1, "vision-fallback")):
        try:
            text = latex_tables_to_markdown(fn(b64)).strip()
            if text:
                return text, name
        except Exception as e:  # noqa: BLE001 - we want to try the fallback on anything
            log.warning("%s failed: %s", name, e)
            errors.append(f"{name}: {e}")
    raise ParseError("; ".join(errors))


def extract_text(filename: str, data: bytes) -> dict:
    """Return {"text", "method", "pages", "warnings"}."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    warnings: list[str] = []

    if ext in {".txt", ".md", ".csv"}:
        return {"text": data.decode("utf-8", errors="replace"), "method": "plain-text", "pages": 1, "warnings": warnings}

    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
        return {"text": "\n".join(parts), "method": "python-docx", "pages": 1, "warnings": warnings}

    if ext in IMAGE_EXT:
        text, model = parse_image(Image.open(io.BytesIO(data)))
        return {"text": text, "method": model, "pages": 1, "warnings": warnings}

    if ext == ".pdf":
        import fitz  # PyMuPDF
        doc = fitz.open(stream=data, filetype="pdf")
        out, methods = [], set()
        for i, page in enumerate(doc):
            if i >= 10:
                warnings.append("Only the first 10 PDF pages were processed.")
                break
            txt = page.get_text().strip()
            if len(txt) > 40:  # real text layer -> no need for a vision model
                out.append(txt)
                methods.add("pdf-text-layer")
            else:               # scanned page -> vision model
                pix = page.get_pixmap(dpi=150)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                try:
                    text, model = parse_image(img)
                    out.append(text)
                    methods.add(model)
                except ParseError as e:
                    warnings.append(f"Page {i + 1} could not be parsed: {e}")
        return {"text": "\n\n".join(out), "method": "+".join(sorted(methods)) or "none",
                "pages": len(doc), "warnings": warnings}

    raise ParseError(f"Unsupported file type '{ext}'. Use txt, md, docx, pdf, png or jpg.")
